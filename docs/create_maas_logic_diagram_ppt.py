from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("docs/maas-methodology-diagram-20260707.pptx")
LEGACY_OUT = Path("docs/maas-logic-diagram-20260707.pptx")
FALLBACK_OUT = Path("docs/maas-methodology-diagram-20260707-agent-folders.pptx")
LATEST_PNG = Path("docs/playwright/design-route-live-verify/maas-20-alt-latest.png")
CRITIQUE_PNG = Path("docs/img_83.png")

NAVY = RGBColor(12, 24, 43)
DARK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
GRAY = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
ORANGE = RGBColor(234, 88, 12)
RED = RGBColor(190, 18, 60)
GREEN = RGBColor(22, 163, 74)
PURPLE = RGBColor(124, 58, 237)
SLATE = RGBColor(51, 65, 85)


def add_text(slide, x, y, w, h, text, size=16, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return box


def add_box(slide, x, y, w, h, text, fill=GRAY, color=DARK, size=12, bold=True, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(203, 213, 225)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Malgun Gothic"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def set_bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def title(slide, text, subtitle=None):
    add_text(slide, 0.45, 0.22, 9.1, 0.42, text, size=22, bold=True, color=DARK)
    if subtitle:
        add_text(slide, 0.47, 0.69, 8.9, 0.32, subtitle, size=10.5, color=MUTED)


def chip(slide, x, y, text, fill, w=1.45):
    return add_box(slide, x, y, w, 0.35, text, fill=fill, color=WHITE, size=9)


def two_col_note(slide, left_title, left_body, right_title, right_body):
    add_box(slide, 0.7, 1.55, 3.9, 0.55, left_title, fill=BLUE, color=WHITE, size=14)
    add_text(slide, 0.85, 2.28, 3.55, 1.65, left_body, size=12, color=DARK)
    add_box(slide, 5.35, 1.55, 3.9, 0.55, right_title, fill=TEAL, color=WHITE, size=14)
    add_text(slide, 5.5, 2.28, 3.55, 1.65, right_body, size=12, color=DARK)


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_text(slide, 0.55, 0.55, 8.9, 0.55, "MAAS: Law-Constrained Architectural Massing", size=27, bold=True, color=WHITE)
    add_text(slide, 0.58, 1.18, 8.7, 0.42, "LLM prompt가 아니라, 언어 표현 + 생성문법 + 법규 solver + agent negotiation + verifier loop", size=14, color=RGBColor(203, 213, 225))
    labels = ["Language\nOntology", "MassDSL\nRepresentation", "Rule\nCompiler", "Legal/Parking\nHard Gates", "Critic +\nRevision"]
    fills = [BLUE, PURPLE, TEAL, ORANGE, GREEN]
    xs = [0.55, 2.35, 4.15, 5.95, 7.75]
    for i, label in enumerate(labels):
        add_box(slide, xs[i], 2.55, 1.45, 0.82, label, fill=fills[i], color=WHITE, size=11)
        if i < len(labels) - 1:
            add_arrow(slide, xs[i] + 1.46, 2.96, xs[i + 1] - 0.04, 2.96, color=WHITE)
    add_text(slide, 0.8, 4.45, 8.4, 0.42, "연구 질문: 법규를 지키면서도 설명 가능한 건축언어 조합과 다양한 mass alternative를 만들 수 있는가?", size=13, bold=True, color=RGBColor(226, 232, 240), align=PP_ALIGN.CENTER)

    # 2. What it is not
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "0. 오해 정리", "LLM에 건축 단어를 넣는 것은 입력 단계일 뿐, 연구 방법론의 끝이 아님")
    two_col_note(
        slide,
        "약한 주장",
        "LLM에게 courtyard, split, roof 같은 단어를 주고\n그럴듯한 도형을 생성한다.\n\n문제: 법규 truth, 주차, repair, 실패 이유,\n재현 가능한 grammar가 없다.",
        "강한 주장",
        "건축언어를 MassDSL로 표현하고,\nrule compiler가 source geometry로 번역하며,\n법규/주차 solver와 critic이 반복 수정한다.\n\n결과는 PNG가 아니라 JSON evidence로도 설명된다.",
    )
    add_box(slide, 2.15, 4.35, 5.7, 0.55, "핵심은 prompt가 아니라 representation + compiler + verifier + revision loop", fill=RED, color=WHITE, size=13)

    # 3. Research lineage
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "1. 논문/참조코드에서 가져온 구조", "정확한 복제보다 구조적 흡수: 표현, 생성, 평가, 반복")
    rows = [
        ("EvoMass", "typology-oriented exploration + population search", "LLM population + island quotas"),
        ("Urban Massing", "massing and layout variation must co-evolve", "parking -> LLM language revision"),
        ("Add/Subtract Massing", "form generation principles create topological variability", "rule_priors + source geometry"),
        ("Agents SDK / GitAgent", "handoffs, tools, traces, repo-native agent identity", "agent.yaml + SOUL/RULES/memory"),
    ]
    y = 1.25
    for name, idea, ours in rows:
        add_box(slide, 0.55, y, 2.0, 0.48, name, fill=SLATE, color=WHITE, size=10)
        add_box(slide, 2.82, y, 3.0, 0.48, idea, fill=GRAY, color=DARK, size=10)
        add_box(slide, 6.1, y, 3.3, 0.48, ours, fill=TEAL, color=WHITE, size=10)
        y += 0.78
    add_text(slide, 0.9, 4.65, 8.2, 0.35, "따라서 발표 핵심은 'LLM이 만들었다'가 아니라 '논문 방법론의 어느 layer가 어느 agent/evidence로 구현됐는가'다.", size=11.5, bold=True, color=DARK, align=PP_ALIGN.CENTER)

    # 4. Representation stack
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "2. 표현 계층", "텍스트 언어가 바로 geometry가 되지 않고, 단계별 contract를 통과")
    levels = [
        ("Architectural\nLanguage", "courtyard, branch, roof, split, interlock", BLUE),
        ("MassDSL\nSequence", "base -> split -> bridge -> void -> roof", PURPLE),
        ("Rule Evidence", "rule_name, authored params, prior params, actions", TEAL),
        ("Source\nGeometry", "volumes, surfaces, primitive roles", ORANGE),
        ("Legalized\nMass", "FAR/BCR/height/parking checked candidate", GREEN),
    ]
    for i, (head, body, fill) in enumerate(levels):
        x = 0.45 + i * 1.9
        add_box(slide, x, 1.35, 1.55, 0.72, head, fill=fill, color=WHITE, size=10)
        add_text(slide, x - 0.05, 2.28, 1.65, 1.05, body, size=9.6, color=DARK, align=PP_ALIGN.CENTER)
        if i < len(levels) - 1:
            add_arrow(slide, x + 1.58, 1.71, x + 1.87, 1.71)
    add_box(slide, 1.6, 4.15, 6.8, 0.52, "각 계층은 JSON evidence로 남아야 하고, 빠진 계층은 verifier failure가 된다.", fill=RED, color=WHITE, size=12)

    # 5. Agent negotiation
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "3. Multi-Agent Constraint Negotiation", "주차 agent와 mass agent 협업은 가능하며, 이 루프가 연구 핵심")
    add_box(slide, 0.35, 1.18, 1.18, 0.62, "Law\nAgent", fill=ORANGE, color=WHITE, size=9.5)
    add_box(slide, 1.72, 1.18, 1.18, 0.62, "Parking\nAgent", fill=TEAL, color=WHITE, size=9.5)
    add_box(slide, 3.09, 1.18, 1.18, 0.62, "LLM\nArchitect", fill=PURPLE, color=WHITE, size=9.5)
    add_box(slide, 4.46, 1.18, 1.18, 0.62, "MassDSL\nAgent", fill=SLATE, color=WHITE, size=9.5)
    add_box(slide, 5.83, 1.18, 1.18, 0.62, "Geometry\nAgent", fill=BLUE, color=WHITE, size=9.5)
    add_box(slide, 7.2, 1.18, 1.18, 0.62, "Critic\nAgent", fill=GREEN, color=WHITE, size=9.5)
    for x1, x2 in [(1.55, 1.70), (2.92, 3.07), (4.29, 4.44), (5.66, 5.81), (7.03, 7.18)]:
        add_arrow(slide, x1, 1.49, x2, 1.49)
    add_arrow(slide, 7.8, 1.9, 2.9, 3.75, color=RED, width=2)
    add_text(slide, 0.8, 2.45, 8.4, 0.48, "예: 주차 실패 -> parking agent가 필로티/기계식/drive aisle 요구 -> LLM architect가 언어 수정 -> MassDSL/geometry가 재컴파일", size=11.5, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    add_box(slide, 1.0, 3.45, 2.0, 0.55, "parking fail", fill=RED, color=WHITE, size=11)
    add_box(slide, 4.0, 3.45, 2.0, 0.55, "revise MassDSL", fill=PURPLE, color=WHITE, size=11)
    add_box(slide, 7.0, 3.45, 2.0, 0.55, "recheck law", fill=ORANGE, color=WHITE, size=11)
    add_arrow(slide, 3.02, 3.72, 3.98, 3.72)
    add_arrow(slide, 6.02, 3.72, 6.98, 3.72)
    add_text(slide, 0.8, 4.7, 8.4, 0.32, "현재 구현은 순차 flow가 강함. 다음 단계는 shared state 기반 bidirectional revision loop.", size=11, color=MUTED, align=PP_ALIGN.CENTER)

    # 6. Agent folder contract
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "4. Agent 폴더 구조와 책임", "GitAgent처럼 agent는 MD 조각이 아니라 독립 폴더 단위")
    add_box(slide, 0.55, 1.05, 2.75, 0.48, "Global Infra", fill=SLATE, color=WHITE, size=12)
    add_text(slide, 0.72, 1.68, 2.45, 1.0, "ARR/backend/agents/\n- A2A discovery\n- JSON-RPC transport\n- worker lifecycle", size=10.2, color=DARK)
    add_box(slide, 3.65, 1.05, 2.75, 0.48, "MAAS Domain", fill=PURPLE, color=WHITE, size=12)
    add_text(slide, 3.82, 1.68, 2.45, 1.0, "ARR/backend/design/maas/agents/\n- law/parking\n- MassDSL/geometry\n- critic/review", size=10.2, color=DARK)
    add_box(slide, 6.75, 1.05, 2.65, 0.48, "Same Folder Contract", fill=TEAL, color=WHITE, size=12)
    add_text(slide, 6.95, 1.62, 2.2, 1.35, "agent.py or adapter\nagent.yaml\nSOUL.md\nRULES.md\nmemory/MEMORY.md\ncard + contract", size=9.6, color=DARK)
    add_arrow(slide, 3.32, 1.29, 3.63, 1.29)
    add_arrow(slide, 6.42, 1.29, 6.73, 1.29)
    add_box(slide, 0.9, 3.32, 2.35, 0.47, "worker module mirror", fill=BLUE, color=WHITE, size=10)
    add_text(slide, 0.72, 3.92, 2.75, 0.48, "worker_agents/modules/{agent_id}/\nkeeps identity/rules/memory", size=9.3, color=DARK, align=PP_ALIGN.CENTER)
    add_box(slide, 3.82, 3.32, 2.35, 0.47, "MAAS specialist package", fill=ORANGE, color=WHITE, size=10)
    add_text(slide, 3.65, 3.92, 2.75, 0.48, "{agent_id}/agent.py + card.py\n+ contract.py + memory", size=9.3, color=DARK, align=PP_ALIGN.CENTER)
    add_box(slide, 6.72, 3.32, 2.35, 0.47, "team graph sync", fill=GREEN, color=WHITE, size=10)
    add_text(slide, 6.55, 3.92, 2.75, 0.48, "JSON_MODULES + AG-light UI\nmust match backend flow", size=9.3, color=DARK, align=PP_ALIGN.CENTER)
    add_box(slide, 1.0, 4.78, 8.0, 0.45, "원칙: global infra와 MAAS domain은 둘 다 필요하지만, 같은 agent-folder 표준으로 검토 가능해야 한다.", fill=RED, color=WHITE, size=11)

    # 7. Rule compiler
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "5. Rule-Based Source Geometry Compiler", "LLM이 낸 언어를 바로 박스로 그리지 않고 rule prior와 authored parameter를 분리")
    add_box(slide, 0.65, 1.25, 2.05, 0.58, "LLM Authored\nParams", fill=PURPLE, color=WHITE, size=11)
    add_box(slide, 3.15, 1.25, 2.05, 0.58, "Architectural\nRule Priors", fill=TEAL, color=WHITE, size=11)
    add_box(slide, 5.65, 1.25, 2.05, 0.58, "Invalid / Missing\nParams", fill=RED, color=WHITE, size=11)
    add_arrow(slide, 2.72, 1.54, 3.12, 1.54)
    add_arrow(slide, 5.22, 1.54, 5.62, 1.54)
    add_text(slide, 0.75, 2.45, 8.4, 0.72, "compiler output = source volumes + source surfaces + rule_evidence + source_signature", size=15, bold=True, color=DARK, align=PP_ALIGN.CENTER)
    chip(slide, 1.05, 3.75, "courtyard", BLUE)
    chip(slide, 2.65, 3.75, "split", BLUE)
    chip(slide, 4.25, 3.75, "interlock", BLUE)
    chip(slide, 5.85, 3.75, "sloped_roof", BLUE, w=1.65)
    chip(slide, 7.65, 3.75, "branch", BLUE)
    add_text(slide, 0.8, 4.65, 8.4, 0.3, "숫자 fallback은 숨기지 않고 rule_prior_params로 evidence화해야 한다.", size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)

    # 8. Legal and parking hard gates
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "6. 법규와 주차는 Hard Gate", "LLM이나 mass agent가 통과 여부를 선언하지 않는다")
    gates = [
        ("FAR", "legal solver", ORANGE),
        ("BCR", "legal solver", ORANGE),
        ("Height", "legal solver", ORANGE),
        ("Sunlight", "envelope/clip", ORANGE),
        ("Parking Count", "parking agent", TEAL),
        ("Mass-Stage Parking", "precheck", TEAL),
    ]
    for i, (name, owner, fill) in enumerate(gates):
        x = 0.75 + (i % 3) * 3.0
        y = 1.25 + (i // 3) * 1.0
        add_box(slide, x, y, 2.25, 0.55, f"{name}\n{owner}", fill=fill, color=WHITE, size=10)
    add_box(slide, 1.1, 3.75, 3.2, 0.55, "통과: legalPass=20/20", fill=GREEN, color=WHITE, size=12)
    add_box(slide, 5.0, 3.75, 3.9, 0.55, "주의: permitParkingPass=0/20", fill=RED, color=WHITE, size=12)
    add_text(slide, 1.0, 4.55, 8.0, 0.34, "현재 주차는 count/mass-stage precheck 통과다. permit-final 주차승인은 별도 swept-path/차로/인허가 검토 필요.", size=10.8, color=MUTED, align=PP_ALIGN.CENTER)

    # 9. Evaluators
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "7. Evaluator가 보는 것", "다양성만 보지 않고 repair와 orderliness까지 동시에 컷")
    evals = [
        ("Diversity", "family, language pair,\nrole diversity", BLUE),
        ("Repair Delta", "legal repair 후\n면적 보존율", ORANGE),
        ("Orderliness", "주축, 위계,\n작은 조각 수", GREEN),
        ("Parking", "required/provided,\nmass-stage", TEAL),
    ]
    for i, (head, body, fill) in enumerate(evals):
        x = 0.65 + i * 2.25
        add_box(slide, x, 1.45, 1.8, 0.58, head, fill=fill, color=WHITE, size=12)
        add_text(slide, x - 0.02, 2.25, 1.84, 0.72, body, size=10, color=DARK, align=PP_ALIGN.CENTER)
    add_box(slide, 1.2, 4.05, 7.6, 0.55, "실패하면 생성기를 약하게 통과시키는 게 아니라, failure memory에 넣고 다음 후보 생성 조건을 바꾼다.", fill=RED, color=WHITE, size=11.5)

    # 10. Current critique from PNG
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "8. 현재 PNG에 대한 정직한 평가", "검증 통과와 건축적 설득력은 다르다")
    if CRITIQUE_PNG.exists():
        slide.shapes.add_picture(str(CRITIQUE_PNG), Inches(0.45), Inches(1.18), width=Inches(4.6))
    elif LATEST_PNG.exists():
        slide.shapes.add_picture(str(LATEST_PNG), Inches(0.45), Inches(1.18), width=Inches(4.6))
    add_box(slide, 5.35, 1.25, 3.95, 0.5, "문제", fill=RED, color=WHITE, size=13)
    add_text(slide, 5.55, 1.95, 3.55, 1.0, "아직 rectilinear box 조합이 강하고,\n일부 후보는 '건축언어'보다 조립 블록처럼 읽힌다.", size=11.2, color=DARK)
    add_box(slide, 5.35, 3.15, 3.95, 0.5, "다음 개선", fill=GREEN, color=WHITE, size=13)
    add_text(slide, 5.55, 3.83, 3.55, 0.95, "agent 협업 loop에서 parking/법규 실패가\nMassDSL language revision으로 돌아가게 만든다.", size=11.2, color=DARK)

    # 11. Revision loop
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "9. Revision Loop", "실패를 prompt 감으로 고치는 게 아니라, 구조화된 action으로 되돌림")
    steps = [
        ("Generate", BLUE),
        ("Compile", PURPLE),
        ("Legalize", ORANGE),
        ("Evaluate", GREEN),
        ("Diagnose", RED),
        ("Revise", TEAL),
    ]
    xs = [0.75, 2.15, 3.55, 4.95, 6.35, 7.75]
    for i, (name, fill) in enumerate(steps):
        add_box(slide, xs[i], 2.0, 1.08, 0.56, name, fill=fill, color=WHITE, size=10)
        if i < len(steps) - 1:
            add_arrow(slide, xs[i] + 1.1, 2.28, xs[i + 1] - 0.02, 2.28)
    add_arrow(slide, 8.3, 2.75, 1.28, 3.95, color=RED, width=2)
    add_text(slide, 0.95, 4.25, 8.1, 0.38, "예: parking fail -> piloti/drive aisle grammar 요청, orderliness low -> fragment 제거/주축 강화 요청", size=11.2, bold=True, color=DARK, align=PP_ALIGN.CENTER)

    # 12. Result metrics
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "10. 현재 검증 결과", "이 숫자는 통과 근거이지, 최종 디자인 품질 보장은 아님")
    metrics = [
        ("legalPass", "20/20", GREEN),
        ("parkingCount", "20/20", GREEN),
        ("parkingPermit", "0/20", RED),
        ("twoTier", "0", GREEN),
        ("orderliness", "0.78 avg", GREEN),
        ("source roles", "57", BLUE),
        ("LLM gap", "0", GREEN),
        ("repair severe", "0", GREEN),
    ]
    for i, (name, value, fill) in enumerate(metrics):
        x = 0.55 + (i % 4) * 2.32
        y = 1.25 + (i // 4) * 0.95
        add_box(slide, x, y, 1.95, 0.58, f"{name}\n{value}", fill=fill, color=WHITE, size=10.5)
    if LATEST_PNG.exists():
        slide.shapes.add_picture(str(LATEST_PNG), Inches(0.85), Inches(3.25), width=Inches(8.3))
    add_text(slide, 0.8, 5.2, 8.4, 0.25, "발표 문장: 현재는 연구용 evidence loop prototype이며, 다음 단계는 agent negotiation 기반 MassDSL revision.", size=9.5, color=MUTED, align=PP_ALIGN.CENTER)

    # 13. Next implementation
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "11. 다음 구현 과제", "지금 부족한 것은 단어 수가 아니라 agent 협업과 geometry 표현력")
    next_items = [
        ("Shared Candidate State", "법규, 주차, MassDSL, geometry, 실패 이유를 한 객체에 저장"),
        ("Parking-Mass Negotiation", "주차 실패가 piloti/void/drive grammar revision으로 환류"),
        ("Non-Box Primitives", "polygonal/curvilinear/freeform source primitive 확대"),
        ("Critic Revision", "PNG/JSON failure를 다음 LLM proposal 조건으로 자동 반영"),
        ("Presentation Evidence", "카드에 rule name, param source, orderliness를 노출"),
    ]
    y = 1.15
    for head, body in next_items:
        add_box(slide, 0.7, y, 2.5, 0.42, head, fill=SLATE, color=WHITE, size=9.8)
        add_text(slide, 3.45, y + 0.04, 5.8, 0.3, body, size=10.5, color=DARK)
        y += 0.72
    add_box(slide, 1.1, 4.85, 7.8, 0.45, "결론: LLM은 제안자, agent loop는 협상자, solver/verifier는 판정자다.", fill=NAVY, color=WHITE, size=12)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    actual_out = OUT
    try:
        prs.save(actual_out)
    except PermissionError:
        actual_out = FALLBACK_OUT
        prs.save(actual_out)
    try:
        copyfile(actual_out, LEGACY_OUT)
    except PermissionError:
        pass
    print(actual_out)
    print(LEGACY_OUT)


if __name__ == "__main__":
    build()
