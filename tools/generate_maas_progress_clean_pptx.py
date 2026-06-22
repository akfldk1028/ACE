from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "maas-progress-report-clean.pptx"
FONT = "Malgun Gothic"

BG = RGBColor(247, 249, 252)
DARK = RGBColor(18, 27, 40)
INK = RGBColor(31, 38, 49)
MUTED = RGBColor(91, 102, 119)
LINE = RGBColor(218, 225, 235)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(28, 99, 156)
GREEN = RGBColor(38, 132, 92)
AMBER = RGBColor(181, 119, 30)
RED = RGBColor(180, 65, 65)


def bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def title(slide, value, sub="MAAS 진행 보고"):
    text(slide, 0.55, 0.32, 2.3, 0.25, sub, 9, BLUE, True)
    text(slide, 0.55, 0.62, 8.9, 0.55, value, 23, INK, True)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.24), Inches(8.9), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def footer(slide, n):
    text(slide, 8.92, 7.1, 0.45, 0.2, f"{n:02d}", 8, MUTED, False, PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, heading, body, color=BLUE):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.055), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    text(slide, x + 0.2, y + 0.14, w - 0.35, 0.25, heading, 11.5, INK, True)
    text(slide, x + 0.2, y + 0.48, w - 0.35, h - 0.55, body, 10.2, MUTED)
    return box


def bullets(slide, x, y, w, h, items, size=12.5, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(4)
    return shape


def picture(slide, rel, x, y, w=None, h=None):
    path = ROOT / rel
    if not path.exists():
        return card(slide, x, y, w or 3, h or 1, "이미지 없음", rel, RED)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)


def table(slide, x, y, w, h, headers, rows, widths=None, size=8.4):
    widths = widths or [1 / len(headers)] * len(headers)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    for i, width in enumerate(widths):
        tbl.columns[i].width = Inches(w * width)
    for c, head in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT
            p.font.size = Pt(size)
            p.font.bold = True
            p.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else RGBColor(241, 245, 249)
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(size)
                p.font.color.rgb = INK
    return shape


def flow_box(slide, x, y, w, h, num, heading, body, color=BLUE):
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE
    text(slide, x + 0.12, y + 0.1, 0.38, 0.28, str(num), 11, color, True)
    text(slide, x + 0.5, y + 0.1, w - 0.6, 0.25, heading, 10.5, INK, True)
    text(slide, x + 0.5, y + 0.38, w - 0.6, h - 0.45, body, 8.6, MUTED)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, DARK)
    text(s, 0.78, 0.82, 8.4, 0.35, "ARR / MAAS", 13, RGBColor(160, 205, 235), True)
    text(s, 0.78, 1.32, 8.5, 1.15, "법규 기반 매스 생성\n진행 보고", 33, WHITE, True)
    text(s, 0.8, 2.8, 8.2, 0.35, "PNU 1168011800104170004 · DOM/API 검증 · VWorld 3D 미검증", 14, RGBColor(221, 230, 240))
    picture(s, "docs/playwright/design-captures/design_optimize_result_with_2d_fallback_1168011800104170004.png", 0.78, 3.62, w=8.45)
    footer(s, 1)

    # 2
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "한 장 요약")
    card(s, 0.7, 1.58, 2.75, 1.25, "핵심", "LLM이 매스를 그리는 구조가 아니다. 실제 geometry는 ARR/MAAS 코드가 만든다.", BLUE)
    card(s, 3.65, 1.58, 2.75, 1.25, "현재 상태", "legal envelope first 방식으로 후보 18개와 evidence bundle을 생성했다.", GREEN)
    card(s, 6.6, 1.58, 2.75, 1.25, "주의", "Playwright 캡처는 VWorld 3D가 아니라 DOM/API와 2D fallback 검증이다.", AMBER)
    bullets(s, 0.92, 3.35, 8.15, 2.25, [
        "발표의 기준: 법규/DEM/datum으로 가능한 공간을 먼저 만들고, 그 안에서 매스 후보를 생성한다.",
        "LLM/agent는 결과를 설명하고 evidence 부족을 지적한다. 좌표나 mesh를 직접 생성하지 않는다.",
        "Nano Banana는 마지막 facade/aesthetic 단계 후보이며, 법규 geometry를 바꾸면 안 된다."
    ], 14)
    footer(s, 2)

    # 3
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "역할 분리")
    card(s, 0.75, 1.55, 2.55, 3.8, "MAAS 코드", "필지 경계, 법규 constraint, legal envelope, floor plates, mass_geojson, repair/check를 계산한다.\n\nGeometry source of truth.", GREEN)
    card(s, 3.72, 1.55, 2.55, 3.8, "LLM / Agent", "후보 설명, evidence 요약, missing evidence 지적, 수정/재생성 요청, review orchestration을 맡는다.\n\nReviewer / explainer.", BLUE)
    card(s, 6.68, 1.55, 2.55, 3.8, "Image Model", "facade, material, window rhythm, presentation image를 만든다.\n\n법규 매스 자체는 변경 금지.", AMBER)
    footer(s, 3)

    # 4
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "메싱 Flow")
    flow_box(s, 0.55, 1.55, 2.75, 0.78, 1, "PNU 입력", "예: 1168011800104170004")
    flow_box(s, 3.62, 1.55, 2.75, 0.78, 2, "대지 정보 수집", "필지, 면적, 용도지역, 도로/인접, DEM datum")
    flow_box(s, 6.7, 1.55, 2.75, 0.78, 3, "법규 constraint", "BCR/FAR/height/setback/일조/채광 reference")
    flow_box(s, 0.55, 2.65, 2.75, 0.78, 4, "legal envelope", "법적으로 매스가 들어갈 수 있는 3D 한계공간")
    flow_box(s, 3.62, 2.65, 2.75, 0.78, 5, "floor plate stack", "층별 허용 바닥판 생성")
    flow_box(s, 6.7, 2.65, 2.75, 0.78, 6, "legal_layered_max", "envelope를 최대한 채우는 기준 후보")
    flow_box(s, 0.55, 3.75, 2.75, 0.78, 7, "seed/operator", "다단 일조스텝, 분절, 테라스, 중정, 바형")
    flow_box(s, 3.62, 3.75, 2.75, 0.78, 8, "repair/check", "법규 밖 후보는 수정하거나 탈락")
    flow_box(s, 6.7, 3.75, 2.75, 0.78, 9, "score/select", "FAR/BCR, 높이, 일조, plan/section diversity")
    flow_box(s, 1.9, 4.85, 2.75, 0.78, 10, "frontend", "mass_geojson, floor_plates, mass_volumes, metrics")
    flow_box(s, 5.35, 4.85, 2.75, 0.78, 11, "agent review", "설명, missing evidence, needs_evidence 판단")
    card(s, 1.0, 6.18, 8.0, 0.55, "정리", "LLM이 좌표를 찍어 매스를 만드는 것이 아니라, 코드가 만든 evidence를 LLM이 검토한다.", RED)
    footer(s, 4)

    # 5
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "법규와 Datum 근거")
    picture(s, "docs/playwright/design-captures/crops/maas_legal_basis_panel.png", 0.65, 1.48, w=3.05)
    card(s, 4.05, 1.52, 2.4, 1.1, "Datum", "NGII local DEM\n대지/도로/인접대지 샘플", GREEN)
    card(s, 6.75, 1.52, 2.4, 1.1, "법규 기준", "§119 대지 기준면\n§86 평균수평면\n정북일조/채광 reference", BLUE)
    card(s, 4.05, 3.0, 2.4, 1.1, "검증값", "도로폭 37.5m\n전면도로 후보 1\n인접대지 후보 4", AMBER)
    card(s, 6.75, 3.0, 2.4, 1.1, "주의", "채광사선은 최종 pass/fail이 아니라 reference. 매스 창면 ray 검토 필요.", RED)
    bullets(s, 4.15, 4.75, 4.8, 0.95, ["법규 수치는 Playwright capture-result-with-fallback.json 화면 텍스트 기준이다."], 12, MUTED)
    footer(s, 5)

    # 6
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Playwright 근거: VWorld가 아니라 2D fallback")
    picture(s, "docs/playwright/design-captures/crops/maas_2d_mass_preview.png", 0.65, 1.45, w=5.25)
    card(s, 6.2, 1.52, 2.85, 1.1, "확인된 것", "DOM 렌더, API flow, 후보 리스트, 2D MASS PREVIEW", GREEN)
    card(s, 6.2, 2.9, 2.85, 1.1, "확인 안 된 것", "Cesium/VWorld 3D 픽셀과 실제 지형 위 배치", RED)
    card(s, 6.2, 4.28, 2.85, 1.1, "그래서 표현", "Playwright DOM/API 캡처로 표기. VWorld 검증 완료라고 쓰지 않음.", AMBER)
    footer(s, 6)

    # 7
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "MAAS 로직 변화")
    card(s, 0.7, 1.55, 2.75, 1.28, "이전 관점", "GA/legacy 후보 생성 → 법규로 걸러냄", RED)
    card(s, 3.65, 1.55, 2.75, 1.28, "현재 관점", "legal envelope → floor plates → 후보 다양화 → repair/check", GREEN)
    card(s, 6.6, 1.55, 2.75, 1.28, "legacy 10종", "capacity source가 아니라 seed/operator source", AMBER)
    bullets(s, 0.9, 3.45, 8.2, 1.8, [
        "legal_layered_max가 기준 후보로 먼저 생성된다.",
        "grammar/morphology 후보는 notch, split, branch, overlap, terrace, tower, taper, grade 등 polygon operator로 만든다.",
        "모든 후보는 repair_design과 legal metrics를 통과해야 UI에 남는다."
    ], 13)
    footer(s, 7)

    # 8
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Seed / Operator 원리")
    table(s, 0.65, 1.55, 8.7, 3.45,
          ["후보", "원리", "실제 연산"],
          [
              ("다단 일조스텝", "층별 법규 envelope 추종", "step_envelope → floor plate stack"),
              ("분절+브릿지+계단", "저층부 분절 + 상부 후퇴", "split → lift → grade"),
              ("테라스", "상부 footprint 축소", "upper scale + volume bands"),
              ("중정형", "중앙 void", "polygon.difference(center box)"),
              ("바형", "한 방향 압축/이동", "scale + translate + intersection"),
              ("코너/오픈코트", "코너 또는 edge 절삭", "difference(cutter box)"),
          ],
          widths=[0.22, 0.31, 0.47],
          size=8.3)
    card(s, 0.9, 5.55, 8.1, 0.72, "핵심", "라벨만 붙인 것이 아니라 Shapely polygon boolean/affine operation 후 legal repair/check를 통과한 후보다.", BLUE)
    footer(s, 8)

    # 9
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "후보 결과")
    table(s, 0.6, 1.52, 8.85, 3.45,
          ["#", "화면명", "operator", "연면적", "일조", "상태"],
          [
              ("1", "법규엔벨로프", "legal_layered_max", "463", "61.0", "OK"),
              ("2", "포디움/타워", "lift_overlap_slabs", "429", "62.4", "OK"),
              ("3", "다단 일조스텝", "grammar_sunlight_multi_step", "463", "61.0", "OK"),
              ("4", "포디움+타워오프셋", "grammar_podium_tower_offset", "309", "61.0", "OK"),
              ("15", "테라스", "grade_terrace", "448", "61.0", "OK"),
              ("18", "분절+브릿지+계단", "grammar_split_lift_stepback", "265", "69.6", "OK"),
          ],
          widths=[0.07, 0.23, 0.35, 0.11, 0.11, 0.13],
          size=8.1)
    card(s, 0.85, 5.35, 3.9, 0.78, "화면 근거", "Playwright capture text에 18개 후보명과 상태가 기록됨.", GREEN)
    card(s, 5.2, 5.35, 3.9, 0.78, "주의", "화면 지표와 evidence dry-run은 산출 시점 차이가 있어 같은 값으로 섞지 않음.", AMBER)
    footer(s, 9)

    # 10
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Evidence Bundle과 Agent Review")
    card(s, 0.75, 1.55, 2.65, 1.25, "상태", "arr.maas.evidence.v0\nchecks 7\nmissing evidence 4\nneeds_evidence", AMBER)
    card(s, 3.68, 1.55, 2.65, 1.25, "담는 것", "site, candidate, geometry, legal, program, assets, provenance, checks", BLUE)
    card(s, 6.62, 1.55, 2.65, 1.25, "Agent 판단", "말로 통과시키지 않고 evidence를 읽어 PASS/FAIL/needs_evidence 판단", GREEN)
    bullets(s, 0.95, 3.45, 8.0, 1.7, [
        "unknown은 pass가 아니다.",
        "geometry는 GeoJSON/assets에 두고, Graph DB는 법규 근거와 lineage provenance에 사용한다.",
        "LLM review는 evidence 위에서만 의미가 있다."
    ], 13)
    footer(s, 10)

    # 11
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "Nano Banana / Aesthetic")
    picture(s, "docs/playwright/design-route-verify/aesthetic-reference/62e70d5c89fd8b59.png", 0.8, 1.58, w=3.0)
    card(s, 4.25, 1.55, 4.85, 1.0, "역할", "legal mass 이후 facade/material/window rhythm을 입히는 presentation 단계", BLUE)
    card(s, 4.25, 2.82, 4.85, 1.0, "구현", "Gemini API 우선 adapter. 기본 모델 gemini-2.5-flash-image-preview.", GREEN)
    card(s, 4.25, 4.1, 4.85, 1.0, "현재 상태", "provider key가 없으면 needs_provider. 생성 성공처럼 표시하지 않음.", AMBER)
    card(s, 4.25, 5.38, 4.85, 0.8, "금지", "silhouette, footprint, height, floor count, setbacks 변경 금지.", RED)
    footer(s, 11)

    # 12
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    title(s, "남은 문제")
    card(s, 0.7, 1.55, 4.15, 1.15, "VWorld", "현재 Playwright는 VWorld 3D 픽셀 검증이 아니다. 실제 브라우저 캡처가 필요하다.", RED)
    card(s, 5.12, 1.55, 4.15, 1.15, "법규", "§119 >3m true area partition, mass-aware daylight, front-road height envelope 미완성.", AMBER)
    card(s, 0.7, 3.0, 4.15, 1.15, "매스", "program/core/corridor/egress/depth rule과 3D diversity validator가 더 필요하다.", AMBER)
    card(s, 5.12, 3.0, 4.15, 1.15, "복잡 필지", "concave/multi-road parcel에서 buildable footprint와 road datum 검증 필요.", RED)
    footer(s, 12)

    # 13
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s, DARK)
    text(s, 0.78, 0.82, 8.5, 0.35, "정리", 15, RGBColor(160, 205, 235), True)
    text(s, 0.78, 1.32, 8.45, 0.95, "MAAS의 핵심은\nLLM 생성이 아니라 검증 가능한 geometry pipeline", 29, WHITE, True)
    card(s, 0.82, 3.0, 2.65, 1.25, "현재 성과", "legal envelope first\n후보 18개\n2D fallback/evidence", GREEN)
    card(s, 3.68, 3.0, 2.65, 1.25, "정확한 표현", "DOM/API 검증\nVWorld 3D 미검증\nneeds_evidence", AMBER)
    card(s, 6.55, 3.0, 2.65, 1.25, "다음", "real-browser VWorld\nmass-aware daylight\nprogram validator", BLUE)
    text(s, 0.84, 5.35, 8.1, 0.42, "산출 파일: docs/maas-progress-report-clean.pptx", 14, RGBColor(220, 230, 240))
    footer(s, 13)

    prs.save(OUT)


if __name__ == "__main__":
    make_deck()
    print(OUT)
