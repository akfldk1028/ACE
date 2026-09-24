# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("ACE_architecture_proposal_white_diagram_v2.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

C_BG = RGBColor(250, 251, 248)
C_GRID = RGBColor(218, 224, 218)
C_GRID2 = RGBColor(233, 237, 232)
C_TEXT = RGBColor(31, 41, 51)
C_MUTED = RGBColor(94, 107, 101)
C_TEAL = RGBColor(15, 76, 58)
C_LIGHT = RGBColor(255, 255, 255)
C_RUST = RGBColor(184, 92, 56)
FONT = "Malgun Gothic"


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C_BG


def add_grid(slide):
    x = 0.22
    while x < 13.12:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(0.22), Inches(x), Inches(7.28)
        )
        line.line.color.rgb = C_GRID2
        line.line.width = Pt(0.45)
        x += 0.23
    y = 0.22
    while y < 7.28:
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(0.22), Inches(y), Inches(13.11), Inches(y)
        )
        line.line.color.rgb = C_GRID2
        line.line.width = Pt(0.45)
        y += 0.23
    rect = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.22), Inches(0.22), Inches(12.89), Inches(7.06)
    )
    rect.fill.background()
    rect.line.color.rgb = C_GRID
    rect.line.width = Pt(1.0)


def text(slide, value, x, y, w, h, size=20, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def bullets(slide, values, x, y, w, h, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, value in enumerate(values):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = value
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = C_TEXT
        p.space_after = Pt(10)
    return box


def footer(slide, n):
    text(slide, "ACE Proposal", 0.58, 7.02, 1.2, 0.18, 8, False, C_MUTED)
    text(slide, f"{n:02d}", 12.35, 7.02, 0.45, 0.18, 8, True, C_MUTED, PP_ALIGN.RIGHT)


def add_slide(n):
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_grid(slide)
    footer(slide, n)
    return slide


def title(slide, section, main, sub=None):
    if section:
        text(slide, section, 0.58, 0.42, 2.5, 0.25, 9, True, C_TEAL)
    text(slide, main, 0.58, 0.68, 11.7, 0.72, 31, True, C_TEXT)
    if sub:
        text(slide, sub, 0.6, 1.36, 11.3, 0.35, 14, False, C_MUTED)
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.58), Inches(1.78), Inches(12.7), Inches(1.78)
    )
    line.line.color.rgb = C_GRID
    line.line.width = Pt(1)


def label(slide, value, x, y, w, h, fill=C_LIGHT, color=C_TEXT, line=C_TEAL, size=13):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line
    box.line.width = Pt(1.1)
    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=C_TEAL, width=1.3, dash=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def iso_mass(slide, x, y, s=1.0):
    def ln(a, b, c, d, color=C_TEAL, width=1.4, dash=False):
        line(slide, x + a * s, y + b * s, x + c * s, y + d * s, color, width, dash)

    base = [(0, 2.0), (2.2, 1.05), (5.0, 2.0), (2.8, 3.08), (0, 2.0)]
    for (a, b), (c, d) in zip(base, base[1:]):
        ln(a, b, c, d, C_TEAL, 1.8)
    blocks = [(1.0, 1.55, 1.2, 1.2), (2.0, 1.25, 1.35, 1.7), (3.0, 1.65, 1.1, 1.0)]
    for bx, by, bw, bz in blocks:
        ln(bx, by, bx + bw, by - 0.35)
        ln(bx + bw, by - 0.35, bx + bw + 0.55, by - 0.12)
        ln(bx + bw + 0.55, by - 0.12, bx + 0.55, by + 0.22)
        ln(bx + 0.55, by + 0.22, bx, by)
        for px, py in [(bx, by), (bx + bw, by - 0.35), (bx + bw + 0.55, by - 0.12), (bx + 0.55, by + 0.22)]:
            ln(px, py, px, py - bz, C_TEAL, 1.1)
        ln(bx, by - bz, bx + bw, by - 0.35 - bz)
        ln(bx + bw, by - 0.35 - bz, bx + bw + 0.55, by - 0.12 - bz)
        ln(bx + bw + 0.55, by - 0.12 - bz, bx + 0.55, by + 0.22 - bz)
        ln(bx + 0.55, by + 0.22 - bz, bx, by - bz)
    for dx in [-0.6, 0.8, 2.5, 4.2, 5.2]:
        ln(dx, 3.55, dx + 2.4, 2.48, C_GRID, 0.75, True)
    for dx in [0.3, 2.1, 4.1]:
        ln(dx, 0.55, dx + 1.8, 1.27, C_GRID, 0.75, True)


def graph(slide, x, y, s=1.0):
    nodes = [
        ("법규", 0, 0),
        ("대지", 1.8, 0.25),
        ("주차", 3.1, 1.15),
        ("후보안", 1.45, 1.65),
        ("위반", -0.15, 1.3),
        ("수정요청", 0.9, 2.55),
        ("근거", 2.75, 2.45),
    ]
    for a, b in [(0, 1), (1, 2), (2, 3), (0, 4), (4, 5), (5, 3), (3, 6), (6, 0)]:
        n1, n2 = nodes[a], nodes[b]
        line(slide, x + n1[1] * s + 0.18, y + n1[2] * s + 0.18, x + n2[1] * s + 0.18, y + n2[2] * s + 0.18, C_GRID, 1.0, True)
    for value, nx, ny in nodes:
        if value in ("후보안", "근거"):
            label(slide, value, x + nx * s, y + ny * s, 0.88 * s, 0.38 * s, C_TEAL, C_LIGHT, C_TEAL, 9)
        else:
            label(slide, value, x + nx * s, y + ny * s, 0.88 * s, 0.38 * s, C_LIGHT, C_TEXT, C_TEAL, 9)


def process(slide, values, x, y):
    for i, value in enumerate(values):
        label(slide, value, x + i * 1.36, y, 1.05, 0.58, C_LIGHT, C_TEXT, C_TEAL, 12)
        if i < len(values) - 1:
            line(slide, x + i * 1.36 + 1.05, y + 0.29, x + (i + 1) * 1.36 - 0.07, y + 0.29, C_TEAL, 1.2)


slides = [
    ("서론", "초기 건축설계의 복잡성", "대지 조건, 용도지역, 프로그램, 법규, 주차, 일조, 도로 조건이 동시에 작동합니다.", ["형태 생성보다 조건 충족과 검토 반복이 핵심", "법규와 geometry가 서로 영향을 주는 구조", "초기 단계에서 오류를 발견하지 못하면 수정 비용 증가"]),
    ("서론", "건축 법규는 하나의 조항이 아닙니다", "법, 시행령, 시행규칙, 주차장법, 조례, 용도지역 조건이 설계 객체 위에 겹쳐집니다.", ["건축법 → 시행령 → 시행규칙", "주차장법 및 지자체 조례", "도로·일조·사선·건폐율·용적률"]),
    ("서론", "주차는 법규와 geometry가 만나는 대표 사례입니다", "필요 주차 대수 산정과 실제 배치 가능성은 서로 다른 문제입니다.", ["법적 산정: 용도와 면적에 따른 필요 대수", "공간 검토: 진입, 회전, 모듈, 기둥, 코어와 충돌", "주차를 통과하지 못하면 설계안 전체가 흔들림"]),
    ("서론", "기존 워크플로우는 생성과 검토가 분리되어 있습니다", "설계안 작성 후 검토하고, 위반 발견 후 다시 수정하는 반복이 발생합니다.", ["설계안 생성 → 법규 검토 → 위반 발견", "수정 후 다른 조건이 다시 깨지는 연쇄", "검토 근거와 수정 이력이 남지 않음"]),
    ("서론", "단일 생성 AI만으로는 부족합니다", "그럴듯한 형태와 법적으로 설명 가능한 설계안은 다릅니다.", ["AI는 시각적으로 plausible한 대안을 만들 수 있음", "하지만 조항 근거, 계산, geometry 검증을 보장하지 못함", "검증 가능한 생성 구조가 필요함"]),
    ("선행 연구", "최신 연구는 네 방향으로 발전 중입니다", "생성, 최적화, 제약조건 처리, 법규 검토 연구가 병렬로 진전되고 있습니다.", ["Optimization-based massing", "Constraint-aware generation", "Architectural 3D / diffusion", "Code compliance / legal AI"]),
    ("선행 연구", "Optimization-Based Massing", "Lou 2025, Wang 2024, EvoMass/SSIEA 계열은 후보안 탐색과 성능 비교를 보여줍니다.", ["강점: 다중 대안 생성과 성능 기반 비교", "한계: 한국 법규 근거와 evidence 추적은 별도 문제", "본 연구: 법규 Graph DB와 검증 루프에 통합"]),
    ("선행 연구", "Constraint-Aware Generation", "PGDM 2024, CDD 2025, PPR 2026은 생성 과정에 제약을 넣으려는 최신 흐름입니다.", ["강점: 생성 모델에 constraint를 결합하려는 시도", "한계: 한국식 사선·일조·주차 등 복합 법규 적용은 미해결", "본 연구: deterministic verifier와 repair loop를 사용"]),
    ("선행 연구", "Architectural 3D / Diffusion", "Tsai 2025, Zhang 2024, UniTEX 2026은 건축 3D와 외피 생성 가능성을 보여줍니다.", ["강점: 시각적 품질과 3D 표현력", "한계: 법규 geometry를 보장하지 않음", "본 연구: 법규 통과 mass 이후 시각화 단계로 연결"]),
    ("선행 연구", "Code Compliance / Legal AI", "LLM-FuncMapper, BuildThemis, LLM-assisted compliance는 법규 해석과 검토 자동화를 다룹니다.", ["강점: 법규 텍스트를 계산/검토로 연결", "한계: 설계안 생성·수정·후보 lineage와 약하게 연결", "본 연구: 법규 검토를 생성 루프 안에 배치"]),
    ("문제 제기", "연구 공백: 통합 루프가 부족합니다", "생성, 검토, 근거, 수정, 협업이 하나의 설계 프로세스로 연결되어야 합니다.", ["형태 생성만으로는 부족", "검토 시스템만으로도 부족", "Evidence와 decision lineage가 필요"]),
    ("문제 제기", "문제 1: 법규 지식이 파편화되어 있습니다", "법령, 시행령, 시행규칙, 조례, 대지 조건이 서로 다른 층위에 존재합니다.", ["문서 중심 저장은 관계 추적에 약함", "어떤 조항이 어떤 후보안에 영향을 주는지 연결 필요", "Graph DB 기반 구조화가 필요"]),
    ("문제 제기", "문제 2: 법규 조건은 geometry로 변환됩니다", "건폐율, 용적률, 사선, 일조, 주차는 텍스트가 아니라 공간 조건으로 작동합니다.", ["법규 해석만으로는 충분하지 않음", "수치 계산과 3D/2D geometry 검증이 필요", "deterministic validator가 기준이 되어야 함"]),
    ("문제 제기", "문제 3: 협업 없는 AI는 책임 경계가 흐립니다", "법규 검토, 주차 검토, 형태 수정, 최종 판단은 성격이 다른 작업입니다.", ["단일 agent는 오류 원인을 분리하기 어려움", "역할별 agent가 evidence를 남겨야 함", "AI는 판단을 보조하고 검증은 도구가 수행"]),
    ("연구 설계", "제안 방법 개요", "Graph DB, MAAS 생성기, deterministic validator, multi-agent review를 하나의 루프로 연결합니다.", ["PNU/대지 입력", "법규 Graph DB", "MAAS 후보안 생성", "법규·주차·형태 검증", "agent repair/review", "evidence 저장"]),
    ("연구 설계", "Graph DB Evidence Model", "법규, 대지, 후보안, 위반, 수정요청, evidence, decision을 관계로 저장합니다.", ["Parcel → Law Article → Constraint", "Candidate → Violation → Repair Request", "Agent Review → Evidence Artifact → Decision"]),
    ("연구 설계", "Multi-Agent Collaboration", "에이전트는 geometry truth를 만들지 않고, 검증 도구와 evidence를 중심으로 협업합니다.", ["Law Agent: 조항과 근거", "Parking Agent: 대수와 배치 가능성", "MAAS/Geometry Agent: 후보 생성과 수정", "Review Agent: 통과/실패와 근거 요약"]),
    ("연구 설계", "Generate–Verify–Repair Loop", "후보안을 만들고, 검증하고, 위반을 repair request로 바꾸어 다시 생성합니다.", ["생성된 후보는 즉시 검증 대상", "위반은 탈락 또는 수정 요청으로 기록", "후보 lineage를 남겨 최종 선택 근거 확보"]),
    ("연구 설계", "Parking Evidence Loop", "주차는 법규 계산과 실제 배치 가능성을 함께 검증하는 핵심 사례입니다.", ["법적 필요 대수 산정", "layout feasibility 검토", "실패 시 core/void/footprint repair 요청"]),
    ("연구 설계", "구현 현황", "현재 시스템은 MAAS 후보안 생성, section profile, parking graph, AG-light 협업 구조를 포함합니다.", ["MAAS benchmark 및 section-profile verification", "parking graph/count/layout checks", "Neo4j law/provenance graph", "AG-light multi-agent team"]),
    ("연구 설계", "검증 지표", "현재 검증 결과는 후보 다양성과 법규 통과, section connector evidence를 중심으로 기록됩니다.", ["successful_scenarios = 4/4", "feature_count = 80", "scenario당 20 unique shapes", "section_connector_feature_count = 11", "legal_pass_rate = 1.0"]),
    ("연구 설계", "평가 계획", "생성 품질보다 “검증 가능한 설계안 생성”을 중심으로 평가합니다.", ["법규 통과율과 위반 감소", "주차 배치 가능성", "후보 다양성", "evidence completeness", "agent repair trace"]),
    ("결론", "기대 효과", "초기 설계 단계에서 법규 위반을 조기에 발견하고, 검토 근거를 추적할 수 있습니다.", ["설계안 생성과 검토의 통합", "법규 검토 근거의 시각화", "후보안 수정 이력과 최종 decision 추적", "설계자-AI 협업 가능성"]),
    ("결론", "연구 기여", "본 연구의 기여는 알고리즘 발명이 아니라 건축설계 검증 프로세스의 통합 구조입니다.", ["한국 건축 법규 Graph DB 구조", "deterministic verifier 기반 생성-검증 루프", "multi-agent 협업형 repair/review workflow", "evidence와 candidate lineage 모델"]),
    ("결론", "결론", "건축설계 자동화의 목표는 단순 생성이 아니라 검증 가능한 생성입니다.", ["법규와 geometry를 함께 다루는 설계 자동화", "Graph DB 기반 근거 추적", "협업형 multi-agent 설계 검증", "초기 설계안 생성·검토 프로세스 확장"]),
]


s = add_slide(1)
text(s, "Graph DB와 협업형\n멀티에이전트를 활용한\n건축설계 자동화 및\n법규 검증 방법 연구", 0.65, 1.12, 6.1, 2.7, 31, True)
text(s, "대지·법규 지식 구조화와 에이전트 협업을 통한\n초기 설계안 생성 및 검토 프로세스", 0.68, 4.35, 6.0, 0.75, 17)
iso_mass(s, 6.45, 1.2, 1.05)
graph(s, 6.15, 0.75, 0.78)

s = add_slide(2)
title(s, "", "건축 초기 설계는 조건 충족과 의사결정의 반복입니다")
for i, item in enumerate(["1. 서론", "2. 선행 연구", "3. 문제 제기", "4. 연구 설계 방향", "5. 결론"]):
    text(s, item, 1.1, 2.05 + i * 0.72, 3.5, 0.35, 20, i == 0)
process(s, ["법규", "대지", "주차", "일조", "용적률", "형태"], 5.3, 2.25)
iso_mass(s, 7.1, 3.55, 0.62)

for idx, (sec, main, sub, items) in enumerate(slides, start=3):
    s = add_slide(idx)
    title(s, sec, main, sub)
    bullets(s, items, 0.8, 2.25, 5.55, 3.85, 18)
    if idx in (7, 8, 12, 13, 14, 16, 17, 18, 19, 20, 21):
        graph(s, 7.0, 2.05, 0.9)
    else:
        iso_mass(s, 7.1, 2.1, 0.72)
        if idx in (15, 22, 23, 24):
            process(s, ["입력", "생성", "검증", "수정", "근거"], 6.55, 5.75)

prs.save(OUT)
print(OUT)
